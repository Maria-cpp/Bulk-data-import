"""Word document table extraction using python-docx."""

import io
from typing import Any

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.services.extraction.base import (
    BaseExtractor,
    ExtractionResult,
    ExtractedTableData,
    ExtractionWarning,
)


class WordExtractor(BaseExtractor):
    """Extract tables from Word documents using python-docx."""

    @property
    def supported_mime_types(self) -> list[str]:
        return [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ]

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        """Extract tables from Word document."""
        tables: list[ExtractedTableData] = []

        try:
            doc = Document(io.BytesIO(file_bytes))

            if not doc.tables:
                return self._create_error_result(
                    "No tables found in Word document",
                    {"filename": filename, "paragraphs_count": len(doc.paragraphs)},
                )

            for table_index, table in enumerate(doc.tables):
                extracted = self._process_table(table, table_index, filename)
                if extracted:
                    tables.append(extracted)

        except PackageNotFoundError:
            return self._create_error_result(
                "Invalid or corrupted Word document",
                {"filename": filename},
            )
        except Exception as e:
            return self._create_error_result(
                f"Failed to process Word document: {str(e)}",
                {"filename": filename, "error_type": type(e).__name__},
                partial_tables=tables if tables else None,
            )

        if not tables:
            return self._create_error_result(
                "No valid tables could be extracted from Word document",
                {"filename": filename},
            )

        return self._create_success_result(tables)

    def _process_table(
        self,
        table: Any,
        table_index: int,
        filename: str,
    ) -> ExtractedTableData | None:
        """Process a single Word table into ExtractedTableData."""
        warnings: list[ExtractionWarning] = []

        rows = table.rows
        if not rows:
            return None

        # First row as headers
        header_row = rows[0]
        raw_headers = [self._get_cell_text(cell) for cell in header_row.cells]

        # Handle merged cells - deduplicate consecutive identical headers
        headers = self._deduplicate_headers(raw_headers, warnings, table_index)

        if not headers:
            warnings.append(
                ExtractionWarning(
                    type="empty_headers",
                    location=f"Table {table_index + 1}",
                    message="Table has no valid headers",
                )
            )
            return None

        # Process data rows
        row_data = []
        for row_idx, row in enumerate(rows[1:], start=1):
            cells = [self._get_cell_text(cell) for cell in row.cells]

            # Handle merged cells - deduplicate values matching header pattern
            if len(cells) != len(headers):
                cells = self._align_cells_to_headers(cells, raw_headers, headers)

            row_dict = {}
            for col_idx, header in enumerate(headers):
                if col_idx < len(cells):
                    value = cells[col_idx]
                    row_dict[header] = value if value else None
                else:
                    row_dict[header] = None

            # Skip completely empty rows
            if any(v is not None and v != "" for v in row_dict.values()):
                row_data.append(row_dict)

        if not row_data:
            warnings.append(
                ExtractionWarning(
                    type="empty_data",
                    location=f"Table {table_index + 1}",
                    message="Table has headers but no data rows",
                )
            )

        return ExtractedTableData(
            table_index=table_index,
            source_location=f"Table {table_index + 1}",
            column_headers=headers,
            row_data=row_data,
            row_count=len(row_data),
            column_count=len(headers),
            warnings=warnings,
        )

    def _get_cell_text(self, cell: Any) -> str:
        """Extract text from a Word table cell."""
        text = cell.text.strip()
        # Clean up whitespace
        text = " ".join(text.split())
        return text

    def _deduplicate_headers(
        self,
        raw_headers: list[str],
        warnings: list[ExtractionWarning],
        table_index: int,
    ) -> list[str]:
        """Deduplicate headers, handling merged cells."""
        headers = []
        seen = {}
        prev_header = None

        for idx, header in enumerate(raw_headers):
            # Skip consecutive duplicates (merged cells)
            if header == prev_header and header:
                continue

            prev_header = header

            h = header if header else f"Column_{len(headers) + 1}"

            if not header:
                warnings.append(
                    ExtractionWarning(
                        type="missing_header",
                        location=f"Table {table_index + 1}, Column {idx + 1}",
                        message=f"Empty header replaced with '{h}'",
                    )
                )

            # Handle remaining duplicates
            if h in seen:
                seen[h] += 1
                original = h
                h = f"{h}_{seen[h]}"
                warnings.append(
                    ExtractionWarning(
                        type="duplicate_header",
                        location=f"Table {table_index + 1}, Column {idx + 1}",
                        message=f"Duplicate header '{original}' renamed to '{h}'",
                    )
                )
            else:
                seen[h] = 0

            headers.append(h)

        return headers

    def _align_cells_to_headers(
        self,
        cells: list[str],
        raw_headers: list[str],
        headers: list[str],
    ) -> list[str]:
        """Align cell values to deduplicated headers, handling merged cells."""
        if len(cells) == len(raw_headers):
            # Same count as raw headers, need to merge cells matching merged headers
            aligned = []
            prev_header = None
            prev_idx = -1

            for idx, header in enumerate(raw_headers):
                if header == prev_header and header:
                    # Merged cell - skip (value is same as previous)
                    continue
                prev_header = header
                prev_idx = idx
                if idx < len(cells):
                    aligned.append(cells[idx])

            return aligned

        # Otherwise, try to match by position
        return cells[:len(headers)]
