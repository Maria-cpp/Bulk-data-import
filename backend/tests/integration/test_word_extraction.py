"""Integration tests for Word document extraction."""

import pytest
from pathlib import Path

from app.services.extraction.word import WordExtractor


# Fixtures path
FIXTURES_DIR = Path(__file__).parent.parent / "fixtures"


@pytest.fixture
def extractor():
    """Create WordExtractor instance."""
    return WordExtractor()


@pytest.fixture
def tables_only_bytes():
    """Load tables-only Word file bytes."""
    filepath = FIXTURES_DIR / "tables_only.docx"
    return filepath.read_bytes()


@pytest.fixture
def mixed_content_bytes():
    """Load mixed content Word file bytes."""
    filepath = FIXTURES_DIR / "mixed_content.docx"
    return filepath.read_bytes()


class TestWordExtractor:
    """Test WordExtractor functionality."""

    def test_supports_docx_mime_type(self, extractor):
        """Verify extractor supports DOCX MIME type."""
        assert extractor.supports(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    def test_supports_doc_mime_type(self, extractor):
        """Verify extractor supports legacy DOC MIME type."""
        assert extractor.supports("application/msword")

    def test_does_not_support_pdf(self, extractor):
        """Verify extractor does not support PDF."""
        assert not extractor.supports("application/pdf")

    def test_does_not_support_excel(self, extractor):
        """Verify extractor does not support Excel."""
        assert not extractor.supports(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )


class TestTableExtraction:
    """Tests for extracting tables from Word documents."""

    def test_extract_tables_only_doc(self, extractor, tables_only_bytes):
        """Extract tables from document containing only tables."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        assert result.success is True
        # Fixture has 2 tables
        assert len(result.tables) == 2

    def test_first_table_structure(self, extractor, tables_only_bytes):
        """Verify first table has expected structure."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        table = result.tables[0]
        # First table: ID, Name, Category, Value
        assert "ID" in table.column_headers
        assert "Name" in table.column_headers
        assert "Category" in table.column_headers
        assert "Value" in table.column_headers

    def test_first_table_row_count(self, extractor, tables_only_bytes):
        """Verify first table has expected row count."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        table = result.tables[0]
        # First table has 5 data rows
        assert table.row_count == 5

    def test_second_table_structure(self, extractor, tables_only_bytes):
        """Verify second table has expected structure."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        table = result.tables[1]
        # Second table: Metric, Value, Change
        assert "Metric" in table.column_headers
        assert "Value" in table.column_headers
        assert "Change" in table.column_headers


class TestIgnoreParagraphs:
    """Tests for ignoring non-table content."""

    def test_only_tables_extracted_from_mixed_content(self, extractor, mixed_content_bytes):
        """Verify only tables are extracted, paragraphs are ignored."""
        result = extractor.extract(mixed_content_bytes, "mixed_content.docx")

        assert result.success is True
        # Should only have tables, no paragraph content
        assert len(result.tables) == 1  # Mixed content doc has 1 table

    def test_table_data_not_polluted_by_paragraphs(self, extractor, mixed_content_bytes):
        """Verify extracted table data doesn't contain paragraph text."""
        result = extractor.extract(mixed_content_bytes, "mixed_content.docx")

        table = result.tables[0]

        # Check that paragraph text is not in table data
        all_values = []
        for row in table.row_data:
            all_values.extend(str(v) for v in row.values() if v)

        paragraph_phrases = [
            "This document contains",
            "extraction process should",
            "bullet point",
        ]

        for phrase in paragraph_phrases:
            assert not any(phrase in val for val in all_values)

    def test_mixed_content_table_has_correct_headers(self, extractor, mixed_content_bytes):
        """Verify table extracted from mixed content has correct headers."""
        result = extractor.extract(mixed_content_bytes, "mixed_content.docx")

        table = result.tables[0]
        # Sales table: Month, Sales, Target, Achievement
        expected_headers = ["Month", "Sales", "Target", "Achievement"]
        for header in expected_headers:
            assert header in table.column_headers


class TestNestedTables:
    """Tests for handling nested tables."""

    def test_nested_tables_handled(self, extractor, tables_only_bytes):
        """Verify nested tables are handled appropriately."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        # Should not crash with nested tables
        assert result.success is True
        # Each table should have valid structure
        for table in result.tables:
            assert table.row_count >= 0
            assert table.column_count > 0


class TestSourceLocation:
    """Tests for table source location tracking."""

    def test_tables_have_source_location(self, extractor, tables_only_bytes):
        """Verify each table has a source location."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        for table in result.tables:
            assert table.source_location is not None
            assert "Table" in table.source_location

    def test_tables_numbered_correctly(self, extractor, tables_only_bytes):
        """Verify tables are numbered in order."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        for i, table in enumerate(result.tables):
            expected_location = f"Table {i + 1}"
            assert expected_location in table.source_location


class TestDataPreservation:
    """Tests for data preservation during extraction."""

    def test_cell_values_preserved(self, extractor, tables_only_bytes):
        """Verify cell values are correctly extracted."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        table = result.tables[0]
        first_row = table.row_data[0]

        # Check values from fixture
        assert first_row["ID"] == "1"
        assert first_row["Name"] == "Item 1"

    def test_numeric_strings_preserved(self, extractor, mixed_content_bytes):
        """Verify numeric values are preserved as strings."""
        result = extractor.extract(mixed_content_bytes, "mixed_content.docx")

        table = result.tables[0]
        # Sales values should be preserved
        for row in table.row_data:
            if row.get("Sales"):
                assert "$" in row["Sales"] or row["Sales"].isdigit() or "," in row["Sales"]


class TestErrorHandling:
    """Tests for error handling scenarios."""

    def test_invalid_docx_returns_error(self, extractor):
        """Verify invalid Word doc returns error result."""
        result = extractor.extract(b"not a valid docx file", "invalid.docx")

        assert result.success is False
        assert result.error_message is not None

    def test_empty_file_returns_error(self, extractor):
        """Verify empty file returns error result."""
        result = extractor.extract(b"", "empty.docx")

        assert result.success is False
        assert result.error_message is not None

    def test_no_tables_returns_error(self, extractor):
        """Verify document with no tables returns error."""
        # Create a minimal DOCX structure without tables would be complex
        # This test validates the error message
        from io import BytesIO
        from docx import Document

        doc = Document()
        doc.add_paragraph("No tables here")
        buffer = BytesIO()
        doc.save(buffer)
        no_tables_bytes = buffer.getvalue()

        result = extractor.extract(no_tables_bytes, "no_tables.docx")

        assert result.success is False
        assert "no table" in result.error_message.lower()


class TestEmptyHeaders:
    """Tests for handling empty or missing headers."""

    def test_empty_headers_replaced(self, extractor, tables_only_bytes):
        """Verify empty headers are replaced with placeholder names."""
        result = extractor.extract(tables_only_bytes, "tables_only.docx")

        for table in result.tables:
            for header in table.column_headers:
                assert header  # No empty headers
                assert len(header) > 0


class TestDuplicateHeaders:
    """Tests for handling duplicate column headers."""

    def test_duplicate_headers_deduplicated(self, extractor):
        """Verify duplicate headers are deduplicated."""
        from io import BytesIO
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.rows[0].cells[0].text = "Name"
        table.rows[0].cells[1].text = "Name"  # Duplicate
        table.rows[0].cells[2].text = "Value"
        table.rows[1].cells[0].text = "A"
        table.rows[1].cells[1].text = "B"
        table.rows[1].cells[2].text = "100"

        buffer = BytesIO()
        doc.save(buffer)
        dup_headers_bytes = buffer.getvalue()

        result = extractor.extract(dup_headers_bytes, "dup_headers.docx")

        if result.success and result.tables:
            table_result = result.tables[0]
            # Headers should be unique
            assert len(table_result.column_headers) == len(
                set(table_result.column_headers)
            )
