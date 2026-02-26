"""Generate test fixture files for integration tests.

Run this script once to create all fixture files:
    python -m tests.fixtures.generate_fixtures

Requires: openpyxl, reportlab, python-docx, Pillow
"""

import csv
import os
from pathlib import Path

# Get fixtures directory
FIXTURES_DIR = Path(__file__).parent


def create_sample_xlsx():
    """Create sample_data.xlsx with multi-sheet test data (3 sheets, 100 rows total)."""
    try:
        from openpyxl import Workbook
    except ImportError:
        print("openpyxl not installed, skipping Excel fixture")
        return

    wb = Workbook()

    # Sheet 1: Employee data (50 rows)
    ws1 = wb.active
    ws1.title = "Employees"
    ws1.append(["ID", "Name", "Department", "Salary", "Hire Date"])
    for i in range(1, 51):
        ws1.append([
            i,
            f"Employee {i}",
            ["Engineering", "Sales", "Marketing", "HR"][i % 4],
            50000 + (i * 1000),
            f"2023-{(i % 12) + 1:02d}-{(i % 28) + 1:02d}"
        ])

    # Sheet 2: Products (30 rows)
    ws2 = wb.create_sheet("Products")
    ws2.append(["SKU", "Product Name", "Category", "Price", "Stock"])
    for i in range(1, 31):
        ws2.append([
            f"SKU-{i:04d}",
            f"Product {i}",
            ["Electronics", "Clothing", "Home", "Sports"][i % 4],
            19.99 + (i * 5),
            100 + (i * 10)
        ])

    # Sheet 3: Orders (20 rows)
    ws3 = wb.create_sheet("Orders")
    ws3.append(["Order ID", "Customer", "Product", "Quantity", "Total", "Status"])
    for i in range(1, 21):
        ws3.append([
            f"ORD-{i:05d}",
            f"Customer {i}",
            f"Product {(i % 30) + 1}",
            i % 10 + 1,
            (i % 10 + 1) * 29.99,
            ["Pending", "Shipped", "Delivered", "Cancelled"][i % 4]
        ])

    output_path = FIXTURES_DIR / "sample_data.xlsx"
    wb.save(output_path)
    print(f"Created: {output_path}")


def create_sample_csv():
    """Create sample_data.csv with test data (50 rows, various data types)."""
    output_path = FIXTURES_DIR / "sample_data.csv"

    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["ID", "Name", "Email", "Age", "Balance", "Active", "Notes"])
        for i in range(1, 51):
            writer.writerow([
                i,
                f"User {i}",
                f"user{i}@example.com",
                20 + (i % 50),
                1000.50 + (i * 100.25),
                "true" if i % 2 == 0 else "false",
                f"Sample notes for user {i}" if i % 3 == 0 else ""
            ])

    print(f"Created: {output_path}")


def create_single_table_pdf():
    """Create single_table.pdf with one clear table."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
        from reportlab.lib import colors
    except ImportError:
        print("reportlab not installed, skipping PDF fixture")
        return

    output_path = FIXTURES_DIR / "single_table.pdf"

    doc = SimpleDocTemplate(str(output_path), pagesize=letter)
    elements = []

    # Create table data
    data = [["ID", "Product", "Category", "Price", "Stock"]]
    for i in range(1, 21):
        data.append([
            str(i),
            f"Product {i}",
            ["Electronics", "Clothing", "Home"][i % 3],
            f"${19.99 + i * 5:.2f}",
            str(100 + i * 10)
        ])

    table = Table(data)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 12),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
        ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ]))

    elements.append(table)
    doc.build(elements)
    print(f"Created: {output_path}")


def create_multi_table_pdf():
    """Create multi_table.pdf with multiple tables across pages."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        print("reportlab not installed, skipping PDF fixture")
        return

    output_path = FIXTURES_DIR / "multi_table.pdf"

    doc = SimpleDocTemplate(str(output_path), pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()

    # Table 1: Sales Data
    elements.append(Paragraph("Sales Report Q1", styles["Heading1"]))
    data1 = [["Month", "Revenue", "Expenses", "Profit"]]
    for month in ["January", "February", "March"]:
        data1.append([month, "$50,000", "$30,000", "$20,000"])

    table1 = Table(data1)
    table1.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.blue),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ]))
    elements.append(table1)
    elements.append(Spacer(1, 30))

    # Table 2: Employee Stats
    elements.append(Paragraph("Employee Statistics", styles["Heading1"]))
    data2 = [["Department", "Employees", "Avg Salary"]]
    for dept in ["Engineering", "Sales", "Marketing", "HR"]:
        data2.append([dept, "25", "$75,000"])

    table2 = Table(data2)
    table2.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.green),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ]))
    elements.append(table2)

    # Page break and Table 3
    elements.append(PageBreak())
    elements.append(Paragraph("Inventory Summary", styles["Heading1"]))
    data3 = [["SKU", "Item", "Quantity", "Location"]]
    for i in range(1, 16):
        data3.append([f"SKU-{i:03d}", f"Item {i}", str(i * 10), f"Warehouse {chr(65 + i % 3)}"])

    table3 = Table(data3)
    table3.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.red),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ]))
    elements.append(table3)

    doc.build(elements)
    print(f"Created: {output_path}")


def create_clear_table_image():
    """Create clear_table.png with clean printed table image."""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        print("Pillow not installed, skipping image fixture")
        return

    # Create image
    width, height = 600, 400
    img = Image.new("RGB", (width, height), color="white")
    draw = ImageDraw.Draw(img)

    # Table data
    headers = ["ID", "Name", "Value"]
    rows = [
        ["1", "Alpha", "100"],
        ["2", "Beta", "200"],
        ["3", "Gamma", "300"],
        ["4", "Delta", "400"],
        ["5", "Epsilon", "500"],
    ]

    # Draw table
    cell_width = 150
    cell_height = 40
    start_x = 75
    start_y = 50

    try:
        font = ImageFont.truetype("arial.ttf", 16)
    except:
        font = ImageFont.load_default()

    # Draw header row
    for col, header in enumerate(headers):
        x = start_x + col * cell_width
        y = start_y
        draw.rectangle([x, y, x + cell_width, y + cell_height], outline="black", fill="lightgray")
        draw.text((x + 10, y + 10), header, fill="black", font=font)

    # Draw data rows
    for row_idx, row in enumerate(rows):
        for col_idx, cell in enumerate(row):
            x = start_x + col_idx * cell_width
            y = start_y + (row_idx + 1) * cell_height
            draw.rectangle([x, y, x + cell_width, y + cell_height], outline="black", fill="white")
            draw.text((x + 10, y + 10), cell, fill="black", font=font)

    output_path = FIXTURES_DIR / "clear_table.png"
    img.save(output_path)
    print(f"Created: {output_path}")


def create_low_quality_image():
    """Create low_quality.jpg with low resolution/skewed table."""
    try:
        from PIL import Image, ImageDraw, ImageFont, ImageFilter
    except ImportError:
        print("Pillow not installed, skipping image fixture")
        return

    # Create base image
    width, height = 400, 300
    img = Image.new("RGB", (width, height), color="lightyellow")
    draw = ImageDraw.Draw(img)

    # Draw somewhat messy table
    headers = ["Col1", "Col2", "Col3"]
    rows = [
        ["A1", "B1", "C1"],
        ["A2", "B2", "C2"],
        ["A3", "B3", "C3"],
    ]

    cell_width = 100
    cell_height = 40
    start_x = 50
    start_y = 30

    try:
        font = ImageFont.truetype("arial.ttf", 14)
    except:
        font = ImageFont.load_default()

    # Draw with slight inconsistencies
    for row_idx, row in enumerate([headers] + rows):
        for col_idx, cell in enumerate(row):
            x = start_x + col_idx * cell_width + (row_idx % 2) * 2  # slight offset
            y = start_y + row_idx * cell_height
            draw.rectangle([x, y, x + cell_width - 2, y + cell_height - 2], outline="gray")
            draw.text((x + 8, y + 8), cell, fill="darkgray", font=font)

    # Add noise and blur for low quality effect
    img = img.rotate(3, expand=True, fillcolor="lightyellow")  # slight rotation
    img = img.filter(ImageFilter.GaussianBlur(radius=0.5))
    img = img.resize((300, 225))  # reduce resolution

    output_path = FIXTURES_DIR / "low_quality.jpg"
    img.save(output_path, quality=60)
    print(f"Created: {output_path}")


def create_tables_only_docx():
    """Create tables_only.docx with multiple tables."""
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError:
        print("python-docx not installed, skipping Word fixture")
        return

    doc = Document()

    # Table 1: Simple data table
    table1 = doc.add_table(rows=6, cols=4)
    table1.style = "Table Grid"
    headers = ["ID", "Name", "Category", "Value"]
    for idx, header in enumerate(headers):
        table1.rows[0].cells[idx].text = header
    for i in range(1, 6):
        table1.rows[i].cells[0].text = str(i)
        table1.rows[i].cells[1].text = f"Item {i}"
        table1.rows[i].cells[2].text = ["A", "B", "C"][i % 3]
        table1.rows[i].cells[3].text = str(i * 100)

    doc.add_paragraph()  # spacing

    # Table 2: Statistics table
    table2 = doc.add_table(rows=4, cols=3)
    table2.style = "Table Grid"
    stats_data = [
        ["Metric", "Value", "Change"],
        ["Revenue", "$1,000,000", "+15%"],
        ["Users", "50,000", "+20%"],
        ["Orders", "10,000", "+10%"],
    ]
    for row_idx, row_data in enumerate(stats_data):
        for col_idx, cell_text in enumerate(row_data):
            table2.rows[row_idx].cells[col_idx].text = cell_text

    output_path = FIXTURES_DIR / "tables_only.docx"
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_mixed_content_docx():
    """Create mixed_content.docx with tables and paragraphs."""
    try:
        from docx import Document
    except ImportError:
        print("python-docx not installed, skipping Word fixture")
        return

    doc = Document()

    doc.add_heading("Quarterly Report", 0)
    doc.add_paragraph(
        "This document contains a mix of text content and tables. "
        "The extraction process should only capture the tabular data below."
    )

    doc.add_heading("Sales Summary", level=1)
    doc.add_paragraph("The following table shows our Q1 sales data:")

    # Add table
    table = doc.add_table(rows=5, cols=4)
    table.style = "Table Grid"
    data = [
        ["Month", "Sales", "Target", "Achievement"],
        ["January", "$50,000", "$45,000", "111%"],
        ["February", "$48,000", "$45,000", "107%"],
        ["March", "$55,000", "$50,000", "110%"],
        ["Total", "$153,000", "$140,000", "109%"],
    ]
    for row_idx, row_data in enumerate(data):
        for col_idx, cell_text in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = cell_text

    doc.add_paragraph()
    doc.add_paragraph(
        "As shown above, we exceeded our targets for all months. "
        "The team's hard work has paid off significantly."
    )

    doc.add_heading("Notes", level=1)
    doc.add_paragraph("- This is a bullet point that should not be extracted")
    doc.add_paragraph("- Neither should this one")

    output_path = FIXTURES_DIR / "mixed_content.docx"
    doc.save(output_path)
    print(f"Created: {output_path}")


def main():
    """Generate all test fixtures."""
    print("Generating test fixtures...")
    print(f"Output directory: {FIXTURES_DIR}")
    print()

    create_sample_xlsx()
    create_sample_csv()
    create_single_table_pdf()
    create_multi_table_pdf()
    create_clear_table_image()
    create_low_quality_image()
    create_tables_only_docx()
    create_mixed_content_docx()

    print()
    print("Done! All fixtures generated.")


if __name__ == "__main__":
    main()
