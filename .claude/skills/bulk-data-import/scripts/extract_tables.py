#!/usr/bin/env python3
"""
Table Extraction Script

Extracts tables from PDF, Excel, Word, and Image files.
Can be run standalone or imported as a module.

Usage:
    python extract_tables.py <input_file> [--output <output_file>] [--format json|csv]

Examples:
    python extract_tables.py assets.pdf --output extracted.json
    python extract_tables.py data.xlsx --format csv --output extracted.csv
    python extract_tables.py scan.png --output ocr_data.json
"""

import sys
import json
import csv
import io
from pathlib import Path
from typing import List, Dict, Any, Optional


def extract_from_pdf(file_path: str) -> List[Dict[str, Any]]:
    """Extract tables from PDF using pdfplumber."""
    try:
        import pdfplumber
    except ImportError:
        print("Error: pdfplumber not installed. Run: pip install pdfplumber")
        sys.exit(1)

    tables = []
    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages):
            page_tables = page.extract_tables()
            for table_idx, table in enumerate(page_tables):
                if not table or len(table) < 2:
                    continue

                headers = [str(h).strip() if h else f"Column_{i}"
                           for i, h in enumerate(table[0])]
                rows = []
                for row in table[1:]:
                    row_data = {headers[i]: str(row[i]).strip() if i < len(row) and row[i] else ""
                                for i in range(len(headers))}
                    rows.append(row_data)

                tables.append({
                    "source": f"Page {page_num + 1}, Table {table_idx + 1}",
                    "headers": headers,
                    "rows": rows
                })

    return tables


def extract_from_excel(file_path: str) -> List[Dict[str, Any]]:
    """Extract tables from Excel using openpyxl."""
    try:
        import openpyxl
    except ImportError:
        print("Error: openpyxl not installed. Run: pip install openpyxl")
        sys.exit(1)

    tables = []
    wb = openpyxl.load_workbook(file_path, data_only=True)

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        data = list(sheet.values)

        if not data or len(data) < 2:
            continue

        headers = [str(h).strip() if h else f"Column_{i}"
                   for i, h in enumerate(data[0])]
        rows = []
        for row in data[1:]:
            if not any(cell for cell in row):
                continue
            row_data = {headers[i]: str(row[i]).strip() if i < len(row) and row[i] else ""
                        for i in range(len(headers))}
            rows.append(row_data)

        if rows:
            tables.append({
                "source": f"Sheet: {sheet_name}",
                "headers": headers,
                "rows": rows
            })

    return tables


def extract_from_word(file_path: str) -> List[Dict[str, Any]]:
    """Extract tables from Word document using python-docx."""
    try:
        from docx import Document
    except ImportError:
        print("Error: python-docx not installed. Run: pip install python-docx")
        sys.exit(1)

    tables = []
    doc = Document(file_path)

    for table_idx, table in enumerate(doc.tables):
        if len(table.rows) < 2:
            continue

        headers = [cell.text.strip() if cell.text else f"Column_{i}"
                   for i, cell in enumerate(table.rows[0].cells)]
        rows = []
        for row in table.rows[1:]:
            row_data = {headers[i]: row.cells[i].text.strip() if i < len(row.cells) else ""
                        for i in range(len(headers))}
            if any(v for v in row_data.values()):
                rows.append(row_data)

        if rows:
            tables.append({
                "source": f"Table {table_idx + 1}",
                "headers": headers,
                "rows": rows
            })

    return tables


def extract_from_image(file_path: str) -> List[Dict[str, Any]]:
    """Extract text/tables from image using OCR."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        print("Error: pytesseract and Pillow required. Run: pip install pytesseract Pillow")
        sys.exit(1)

    image = Image.open(file_path)
    ocr_data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DATAFRAME)
    ocr_data = ocr_data[ocr_data['text'].notna() & (ocr_data['text'].str.strip() != '')]

    if ocr_data.empty:
        return []

    grouped = ocr_data.groupby(['block_num', 'line_num'])
    extracted_rows = []

    for (block, line), group in grouped:
        group_sorted = group.sort_values('left')
        row_text = group_sorted['text'].tolist()
        if row_text:
            extracted_rows.append(row_text)

    if len(extracted_rows) < 2:
        return []

    max_cols = max(len(row) for row in extracted_rows)
    normalized_rows = []
    for row in extracted_rows:
        while len(row) < max_cols:
            row.append("")
        normalized_rows.append(row[:max_cols])

    headers = normalized_rows[0]
    rows = [{headers[i]: normalized_rows[row_idx][i] for i in range(len(headers))}
            for row_idx in range(1, len(normalized_rows))]

    return [{
        "source": "OCR Extraction",
        "headers": headers,
        "rows": rows
    }]


def extract_from_csv(file_path: str) -> List[Dict[str, Any]]:
    """Extract data from CSV file."""
    with open(file_path, 'r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        rows = list(reader)

    if not rows:
        return []

    headers = list(rows[0].keys())
    return [{
        "source": "CSV File",
        "headers": headers,
        "rows": rows
    }]


def extract_tables(file_path: str) -> List[Dict[str, Any]]:
    """Extract tables from any supported file type."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == '.pdf':
        return extract_from_pdf(file_path)
    elif ext in {'.xlsx', '.xls'}:
        return extract_from_excel(file_path)
    elif ext in {'.docx', '.doc'}:
        return extract_from_word(file_path)
    elif ext in {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tiff'}:
        return extract_from_image(file_path)
    elif ext == '.csv':
        return extract_from_csv(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def flatten_tables(tables: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Flatten all tables into a single list of rows."""
    all_rows = []
    for table in tables:
        all_rows.extend(table["rows"])
    return all_rows


def to_csv_string(rows: List[Dict[str, Any]]) -> str:
    """Convert rows to CSV string."""
    if not rows:
        return ""

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=rows[0].keys())
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue()


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = None
    output_format = "json"

    # Parse arguments
    i = 2
    while i < len(sys.argv):
        if sys.argv[i] == "--output" and i + 1 < len(sys.argv):
            output_file = sys.argv[i + 1]
            i += 2
        elif sys.argv[i] == "--format" and i + 1 < len(sys.argv):
            output_format = sys.argv[i + 1].lower()
            i += 2
        else:
            i += 1

    # Extract tables
    print(f"Extracting tables from: {input_file}")
    tables = extract_tables(input_file)

    if not tables:
        print("No tables found in the file.")
        sys.exit(0)

    # Flatten rows
    all_rows = flatten_tables(tables)
    print(f"Extracted {len(all_rows)} rows from {len(tables)} table(s)")

    # Output
    if output_format == "csv":
        output_data = to_csv_string(all_rows)
    else:
        output_data = json.dumps(all_rows, indent=2)

    if output_file:
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(output_data)
        print(f"Output written to: {output_file}")
    else:
        print("\n" + output_data)


if __name__ == "__main__":
    main()
